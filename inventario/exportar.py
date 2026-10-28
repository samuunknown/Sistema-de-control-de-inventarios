from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def a_excel(titulo, encabezados, filas):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Reporte'
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='212529')
    for fila in filas:
        ws.append(fila)
    for columna in ws.columns:
        ancho = max(len(str(c.value or '')) for c in columna)
        ws.column_dimensions[columna[0].column_letter].width = min(ancho + 2, 45)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _miles(valor):
    return f'{valor:,}'.replace(',', '.') if isinstance(valor, int) else valor


def a_pdf(titulo, encabezados, filas):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=titulo)
    datos = [encabezados] + [[_miles(v) for v in fila] for fila in filas]
    tabla = Table(datos, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#212529')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f1f3f5')]),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
    ]))
    doc.build([Paragraph(titulo, getSampleStyleSheet()['Title']), Spacer(1, 12), tabla])
    return buffer.getvalue()